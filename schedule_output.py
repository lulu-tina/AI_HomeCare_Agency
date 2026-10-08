"""Pure operational schedule/route views. Rendering and export never call map APIs."""
from io import BytesIO
import hashlib
import json
import math

import pandas as pd


def text(value):
    if value is None or pd.isna(value):
        return ''
    return str(value).strip()


def number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None


def day_string(value):
    value = text(value)
    if not value or value == '單日':
        return value
    parsed = pd.to_datetime(value, errors='coerce')
    return parsed.strftime('%Y-%m-%d') if pd.notna(parsed) else value


def clock(value):
    value = text(value)
    if not value:
        return ''
    parsed = pd.to_datetime(value, errors='coerce')
    return parsed.strftime('%H:%M') if pd.notna(parsed) else value


def describe_routes(assignments, tasks, caregivers, buffer_mins=15):
    """Describe home→first visit→subsequent visits, resetting at each worker/day."""
    columns = ['任務ID', '派單居服員', '交通日期', '交通方式', '出發地類型', '起點案號',
               '前一任務ID', '交通起點', '交通終點', '交通路段', '出發地緯度', '出發地經度',
               '終點緯度', '終點經度', '前案結束時間', '交通查詢方式', '交通查詢時間',
               '路段緩衝(分)', '交通資料簽章']
    if assignments.empty:
        return pd.DataFrame(columns=columns)
    task_map = tasks.drop_duplicates('任務ID', keep='last').set_index('任務ID').to_dict('index')
    cg_map = caregivers.drop_duplicates('居服員ID', keep='last').set_index('居服員ID').to_dict('index')
    groups = {}
    for _, row in assignments.drop_duplicates('任務ID', keep='last').iterrows():
        task_id, cg_id = text(row.get('任務ID')), text(row.get('派單居服員'))
        if cg_id in ('', '未指派') or task_id not in task_map:
            continue
        task = task_map[task_id]
        day = day_string(task.get('日期')) or day_string(row.get('排班日期')) or '單日'
        groups.setdefault((cg_id, day), []).append(task_id)
    rows = []
    for (cg_id, day), task_ids in sorted(groups.items()):
        cg = cg_map.get(cg_id, {})
        mode = text(cg.get('常用交通工具')) or text(cg.get('交通工具')) or '未提供'
        if any(word in mode for word in ('公車', '捷運', '大眾')):
            mode = '大眾運輸'
        prev_id = ''
        for task_id in sorted(task_ids, key=lambda key: (clock(task_map[key].get('時間窗_開始')), key)):
            task = task_map[task_id]
            prev = task_map.get(prev_id, {})
            first = not prev_id
            origin_id = text(prev.get('案家ID'))
            origin = '居服員住家／服務起點' if first else f'前案｜{origin_id}（{prev_id}）'
            destination = f"{text(task.get('案家ID'))}（{task_id}）"
            lat1 = number(cg.get('服務起點_緯度(家)')) if first else number(prev.get('服務地點_緯度'))
            lon1 = number(cg.get('服務起點_經度(家)')) if first else number(prev.get('服務地點_經度'))
            lat2, lon2 = number(task.get('服務地點_緯度')), number(task.get('服務地點_經度'))
            buffer = 0.0 if first else float(buffer_mins)
            query_clock = clock(task.get('時間窗_開始')) if first else clock(prev.get('時間窗_結束'))
            query_time = ''
            if day != '單日' and query_clock:
                ts = pd.to_datetime(day + ' ' + query_clock, errors='coerce')
                if pd.notna(ts):
                    if not first:
                        ts += pd.Timedelta(minutes=buffer)
                    query_time = ts.tz_localize('Asia/Taipei').isoformat()
            fingerprint = [task_id, cg_id, day, clock(task.get('時間窗_開始')),
                           clock(task.get('時間窗_結束')), mode, prev_id, query_time,
                           lat1, lon1, lat2, lon2, buffer]
            rows.append(dict(zip(columns, [task_id, cg_id, day, mode,
                '住家／服務起點' if first else '前一案家', origin_id, prev_id,
                origin, destination, origin + ' → ' + destination, lat1, lon1, lat2, lon2,
                clock(prev.get('時間窗_結束')), '到達' if first else '出發', query_time, buffer,
                hashlib.sha256(json.dumps(fingerprint, ensure_ascii=False).encode()).hexdigest()])))
            prev_id = task_id
    return pd.DataFrame(rows, columns=columns)


MAIN_COLUMNS = ['日期', '星期', '開始時間', '結束時間', '案號', '派單居服員', '服務項目',
                '服務分鐘', '出發地', '交通方式', '交通時間', '路段緩衝(分)',
                '含緩衝交通時間(分)', '交通計算狀態', '接單狀態', '任務ID']


def build_schedule_output(tasks, assignments, caregivers, routes, metadata=None, buffer_mins=15):
    """Use only travel calculated for the current worker, predecessor, and time."""
    sequence = describe_routes(assignments, tasks, caregivers, buffer_mins)
    sequence_map = sequence.set_index('任務ID').to_dict('index') if not sequence.empty else {}
    route_map = routes.drop_duplicates('任務ID', keep='last').set_index('任務ID').to_dict('index') if not routes.empty else {}
    meta_map = metadata.drop_duplicates('任務ID', keep='last').set_index('任務ID').to_dict('index') if metadata is not None and not metadata.empty else {}
    rows = []
    for _, task in tasks.drop_duplicates('任務ID', keep='last').iterrows():
        task_id = text(task['任務ID'])
        row = dict(task)
        row.update(meta_map.get(task_id, {}))
        seq = sequence_map.get(task_id)
        route = route_map.get(task_id, {})
        row.update({'任務ID': task_id, '案號': task.get('案家ID', task.get('案號')),
                    '日期': day_string(task.get('日期', '')), '星期': task.get('星期', ''),
                    '開始時間': clock(task.get('時間窗_開始', task.get('開始時間'))),
                    '結束時間': clock(task.get('時間窗_結束', task.get('結束時間')))})
        minutes = number(task.get('服務歷時(分鐘)', task.get('服務分鐘')))
        row['服務分鐘'] = minutes
        row['服務時數'] = round(minutes / 60, 4) if minutes is not None else None
        service_parts = []
        for col in sorted(task.index):
            if col.startswith('Service_Code_') and text(task[col]):
                suffix = col.rsplit('_', 1)[1]
                units = text(task.get('Units_' + suffix))
                service_parts.append(text(task[col]) + (f'({units})' if units else ''))
        row['服務項目'] = '、'.join(service_parts) or text(task.get('服務項目'))
        row['派單居服員'] = seq['派單居服員'] if seq else '未指派'
        row['最終派單居服員'] = row['派單居服員']
        row['交通方式'] = seq['交通方式'] if seq else ''
        row['出發地'] = seq['交通起點'] if seq else ''
        row['交通時間'] = None  # The original weekly-plan zero is a placeholder.
        row['預估車程(分)'] = None
        row['含緩衝交通時間(分)'] = None
        row['路段緩衝(分)'] = seq['路段緩衝(分)'] if seq else None
        row['交通計算狀態'] = '已派單，交通待計算' if seq else '未派單，交通尚未計算'
        row['交通估算來源'] = ''
        row['轉場時段檢查'] = ''
        if seq:
            row.update(seq)
            valid = (text(route.get('派單居服員')) == seq['派單居服員'] and
                     text(route.get('交通資料簽章')) == seq['交通資料簽章'])
            if valid:
                travel = number(route.get('預估車程(分)'))
                row['交通時間'] = travel
                row['預估車程(分)'] = travel
                row['含緩衝交通時間(分)'] = number(route.get('含緩衝交通時間(分)'))
                for col in ('交通計算狀態', '交通估算來源', '交通查詢時間', '轉場時段檢查'):
                    row[col] = route.get(col, '')
            elif route:
                row['交通計算狀態'] = '派單或時段已變更，交通待重算'
        rows.append(row)
    out = pd.DataFrame(rows)
    weekly = ['狀態', '費用類型', '頻率', '生效日', '到期日', '服務時數']
    for col in MAIN_COLUMNS + weekly:
        if col not in out:
            out[col] = None
    order = MAIN_COLUMNS + weekly + [col for col in out if col not in MAIN_COLUMNS + weekly]
    return out[order].sort_values(['日期', '開始時間', '派單居服員', '任務ID'], na_position='last').reset_index(drop=True)


def schedule_excel_bytes(schedule):
    """Export familiar weekly-plan columns plus an auditable per-leg worksheet."""
    route_columns = ['日期', '任務ID', '案號', '派單居服員', '開始時間', '結束時間', '交通方式',
                     '出發地類型', '出發地', '起點案號', '前一任務ID', '前案結束時間',
                     '出發地緯度', '出發地經度', '終點緯度', '終點經度', '交通查詢方式',
                     '交通查詢時間', '交通時間', '路段緩衝(分)', '含緩衝交通時間(分)',
                     '交通估算來源', '交通計算狀態', '轉場時段檢查', '交通資料簽章']
    notes = pd.DataFrame([
        ('交通時間', '系統估計分鐘，不含轉場緩衝；原週服務計畫的0是待計算占位值。空白表示未知，不能當作0分鐘。'),
        ('服務時數', '服務分鐘÷60的小數小時；例如75分鐘=1.25小時，不使用原檔HH.MM式數字。'),
        ('出發地', '每人每天第一案由員工資料的住家／服務起點出發；其他由當日有效派單順序的前一案家出發。非實際GPS軌跡。'),
        ('交通方式', '依員工資料選機車或大眾運輸。大眾運輸可能含步行；此欄不能證明實際搭乘公車／捷運。'),
        ('交通查詢方式', '首案以服務開始作為到達條件；後續案以前案結束+轉場緩衝作為出發條件。OSRM不提供即時交通預測。'),
        ('路段緩衝(分)', '獨立預留停車、步行、門禁等時間；首案顯示0，不代表出門與門禁不用時間。'),
        ('交通資料簽章', '用於防止改派／改期／插單後沿用舊路段交通。交通估算仍須人工核對或實測。'),
        ('班表範圍', '含未指派與臨時插單。逐段交通不含最後一案返家。接單狀態另列，指派不代表已接受服務。'),
    ], columns=['欄位', '說明'])
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        schedule.to_excel(writer, sheet_name='班表', index=False)
        schedule.reindex(columns=route_columns).to_excel(writer, sheet_name='逐段交通', index=False)
        notes.to_excel(writer, sheet_name='欄位說明', index=False)
        from openpyxl.styles import Alignment, Font, PatternFill
        for ws in writer.book:
            ws.freeze_panes = 'E2' if ws.title == '班表' else 'A2'
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='176B65')
                cell.font = Font(color='FFFFFF', bold=True)
                cell.alignment = Alignment(wrap_text=True)
            ws.row_dimensions[1].height = 30
            for col in ws.columns:
                heading = str(col[0].value)
                ws.column_dimensions[col[0].column_letter].width = 48 if heading in ('出發地', '交通路段', '說明') else (27 if '時間' in heading or heading == '任務ID' else 20)
                for cell in col[1:]:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '0.00' if any(word in heading for word in ('交通', '時數', '緩衝')) else '0.######'
                    if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                        cell.data_type = 's'
            if ws.title == '欄位說明':
                ws.column_dimensions['B'].width = 110
    return buf.getvalue()
