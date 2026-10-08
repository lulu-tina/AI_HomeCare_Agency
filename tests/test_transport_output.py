import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import unittest
from unittest.mock import patch, Mock
from types import SimpleNamespace
import pandas as pd
import requests
import osrm_client as client
import caregiver_engine as engine
from schedule_output import describe_routes, build_schedule_output, schedule_excel_bytes
from schedule_checks import incremental_travel
from io import BytesIO
from openpyxl import load_workbook

class Updates(unittest.TestCase):
    def setUp(self):
        self.tasks = pd.DataFrame([
            {'任務ID':'T1','案家ID':'A1','日期':'2026-10-06','時間窗_開始':'09:00','時間窗_結束':'10:00','服務歷時(分鐘)':60,'服務地點_緯度':25.1,'服務地點_經度':121.1,'交通時間':0,'Service_Code_1':'BA13'},
            {'任務ID':'T2','案家ID':'A2','日期':'2026-10-06','時間窗_開始':'11:00','時間窗_結束':'12:00','服務歷時(分鐘)':60,'服務地點_緯度':25.2,'服務地點_經度':121.2,'交通時間':0},
            {'任務ID':'T3','案家ID':'A3','日期':'2026-10-07','時間窗_開始':'09:00','時間窗_結束':'10:15','服務歷時(分鐘)':75,'服務地點_緯度':25.3,'服務地點_經度':121.3,'交通時間':0},
            {'任務ID':'T4','案家ID':'A4','日期':'2026-10-06','時間窗_開始':'12:30','時間窗_結束':'13:00','服務歷時(分鐘)':30,'服務地點_緯度':25.4,'服務地點_經度':121.4,'交通時間':0}])
        self.cg = pd.DataFrame([{'居服員ID':'03816','常用交通工具':'機車','服務起點_緯度(家)':25.,'服務起點_經度(家)':121.}, {'居服員ID':'C2','常用交通工具':'大眾運輸','服務起點_緯度(家)':25.5,'服務起點_經度(家)':121.5}])
        self.assign = pd.DataFrame([{'任務ID':t,'派單居服員':'03816'} for t in ['T1','T2','T3']])
        self.config = engine.PipelineConfig(buffer_mins=15)
    def routes(self):
        with patch.object(engine, 'calc_travel_minutes', return_value=10), patch.object(engine, '_travel_source', return_value='測試OSRM'):
            return engine.calculate_schedule_travel(self.assign, self.tasks, self.cg, self.config)
    def test_sequence_and_export(self):
        routes = self.routes()
        out = build_schedule_output(self.tasks,self.assign,self.cg,routes,buffer_mins=15).set_index('任務ID')
        self.assertEqual(out.loc['T2','起點案號'],'A1')
        self.assertEqual(out.loc['T2','交通查詢時間'],'2026-10-06T10:15:00+08:00')
        self.assertEqual(out.loc['T3','出發地類型'],'住家／服務起點')
        self.assertEqual(out.loc['T3','服務時數'],1.25)
        self.assertEqual(out.loc['T1','交通時間'],10)
        self.assertTrue(pd.isna(out.loc['T4','交通時間']))
        book = load_workbook(BytesIO(schedule_excel_bytes(out.reset_index())))
        self.assertEqual(book.sheetnames,['班表','逐段交通','欄位說明'])
        headers=[c.value for c in book['班表'][1]]
        self.assertEqual(book['班表'].cell(2,headers.index('派單居服員')+1).value,'03816')
    def test_timestamp_dates_preserve_taipei_query_time(self):
        tasks=self.tasks.copy();tasks['日期']=pd.to_datetime(tasks['日期'])
        sequence=describe_routes(self.assign,tasks,self.cg).set_index('任務ID')
        self.assertEqual(sequence.loc['T2','交通查詢時間'],'2026-10-06T10:15:00+08:00')
        self.assertEqual(sequence.loc['T2','交通日期'],'2026-10-06')
    def test_override_invalidates_old_travel(self):
        routes=self.routes();new=self.assign.copy();new.loc[new['任務ID'].eq('T1'),'派單居服員']='C2'
        out=build_schedule_output(self.tasks,new,self.cg,routes).set_index('任務ID')
        self.assertTrue(pd.isna(out.loc['T1','交通時間']))
        self.assertEqual(out.loc['T1','交通方式'],'大眾運輸')
        self.assertTrue(pd.isna(out.loc['T2','交通時間']))
        self.assertEqual(out.loc['T2','出發地類型'],'住家／服務起點')
        self.assertEqual(out.loc['T3','交通時間'],10)
    def test_time_and_coordinate_changes_invalidated(self):
        routes=self.routes();tasks=self.tasks.copy();tasks.loc[tasks['任務ID'].eq('T2'),'時間窗_開始']='11:30'
        out=build_schedule_output(tasks,self.assign,self.cg,routes).set_index('任務ID')
        self.assertTrue(pd.isna(out.loc['T2','交通時間']))
        tasks=self.tasks.copy();tasks.loc[tasks['任務ID'].eq('T2'),'服務地點_緯度']=26
        self.assertTrue(pd.isna(build_schedule_output(tasks,self.assign,self.cg,routes).set_index('任務ID').loc['T2','交通時間']))
    def test_insertion_recalculates_only_affected_day(self):
        cache={};calls=[]
        def calculate(*args):
            calls.append(tuple(args[0]['任務ID']));return self.routes_for(*args)
        incremental_travel(self.assign,self.tasks,self.cg,self.config,cache,calculate)
        calls.clear();incremental_travel(self.assign,self.tasks,self.cg,self.config,cache,calculate)
        self.assertEqual(calls,[])
        tasks=self.tasks.copy();tasks.loc[tasks['任務ID'].eq('T4'),['時間窗_開始','時間窗_結束']]=['10:15','10:45']
        assigned=pd.concat([self.assign,pd.DataFrame([{'任務ID':'T4','派單居服員':'03816'}])],ignore_index=True)
        routes=incremental_travel(assigned,tasks,self.cg,self.config,cache,calculate)
        self.assertEqual(len(calls),1)
        self.assertEqual(routes.set_index('任務ID').loc['T2','前一任務ID'],'T4')
    def routes_for(self,*args):
        with patch.object(engine,'calc_travel_minutes',return_value=10),patch.object(engine,'_travel_source',return_value='測試'):
            return engine.calculate_schedule_travel(*args)
    def test_same_coordinates_flagged(self):
        tasks=self.tasks.copy();tasks.loc[tasks['任務ID'].eq('T2'),['服務地點_緯度','服務地點_經度']]=[25.1,121.1]
        routes=self.routes_for(self.assign,tasks,self.cg,self.config)
        self.assertIn('同座標',routes.set_index('任務ID').loc['T2','交通計算狀態'])
    def test_cloud_retry_timeout_and_table_limit(self):
        response=Mock(status_code=200);response.json.return_value={'code':'Ok','routes':[{'duration':13.5}]}
        with patch.object(client,'setting',side_effect=lambda k,d=None:d), patch.object(client.requests,'get',side_effect=[requests.ReadTimeout('x'),response]) as get:
            result=client.request_json('https://example.invalid','/route')
            self.assertEqual(result['code'],'Ok');self.assertEqual(get.call_count,2)
            self.assertEqual(get.call_args.kwargs['timeout'],(5,30))
            self.assertEqual(client.table_limit('https://example.invalid'),40)
            self.assertEqual(client.table_limit('http://127.0.0.1:5001'),90)
    def test_startup_reuses_ready_check(self):
        client._ready.clear()
        with patch.object(client,'request_json',return_value={'code':'Ok','routes':[{'duration':13.5}]}) as request:
            client.ensure_ready('https://example.invalid');client.ensure_ready('https://example.invalid')
            self.assertEqual(request.call_count,1);self.assertTrue(request.call_args.kwargs['startup'])
    def test_matrix_rejects_malformed_response_without_cache_write(self):
        before=engine._OSRM_TRAVEL_TIME_CACHE.copy()
        with patch.object(engine,'ensure_ready'),patch.object(engine,'request_json',return_value={'code':'Ok','durations':[[3]]}):
            with self.assertRaisesRegex(RuntimeError,'尺寸'):
                engine._fetch_osrm_table_chunk([(25,121)],[(25.1,121.1),(25.2,121.2)],3)
        self.assertEqual(before,engine._OSRM_TRAVEL_TIME_CACHE)
    def test_failed_cloud_never_falls_back(self):
        with patch.object(client,'setting',side_effect=lambda k,d=None:d),patch.object(client.requests,'get',side_effect=requests.ReadTimeout('x')) as get:
            with self.assertRaisesRegex(RuntimeError,'未使用汽車或直線'):
                client.request_json('https://example.invalid','/route')
            self.assertEqual(get.call_count,2)

if __name__=='__main__':unittest.main(verbosity=2)
